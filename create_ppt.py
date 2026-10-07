import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_devops_mlops_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette - Sleek Dark Navy & Cyan (Matching Project Dashboard)
    BG_COLOR = RGBColor(15, 23, 42)         # #0F172A (Slate 900)
    CARD_BG = RGBColor(30, 41, 59)          # #1E293B (Slate 800)
    CARD_BORDER = RGBColor(51, 65, 85)      # #334155 (Slate 700)
    ACCENT_CYAN = RGBColor(56, 189, 248)    # #38BDF8 (Sky 400)
    ACCENT_INDIGO = RGBColor(129, 140, 248)  # #818CF8 (Indigo 400)
    ACCENT_GREEN = RGBColor(52, 211, 153)   # #34D399 (Emerald 400)
    TEXT_WHITE = RGBColor(248, 250, 252)    # #F8FAFC
    TEXT_MUTED = RGBColor(148, 163, 184)    # #94A3B8
    TABLE_HEADER_BG = RGBColor(30, 58, 138)

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR

    def add_header(slide, title_text, category="CAPSTONE REVIEW – 1 | DEVOPS & MLOPS PIPELINE"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.3))
        tf_cat = cat_box.text_frame
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_CYAN
        p_cat.font.name = "Arial"

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(11.7), Inches(0.6))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_WHITE
        p_title.font.name = "Arial"

        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.3), Inches(11.733), Inches(0.03))
        line.fill.solid()
        line.fill.fore_color.rgb = CARD_BORDER
        line.line.fill.background()

    def add_card(slide, left, top, width, height, bg_rgb=CARD_BG, border_rgb=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_rgb
        shape.line.color.rgb = border_rgb
        shape.line.width = Pt(1)
        return shape

    # ==========================================
    # SLIDE 1: Title & Team Information
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)

    add_card(slide1, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9), CARD_BG, ACCENT_CYAN)

    tag_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.1), Inches(9), Inches(0.4))
    p = tag_box.text_frame.paragraphs[0]
    p.text = "CAPSTONE REVIEW – 1 | DEVOPS & MLOPS ORCHESTRATION PIPELINE"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    t_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(10.9), Inches(1.2))
    t_frame = t_box.text_frame
    t_frame.word_wrap = True
    p = t_frame.paragraphs[0]
    p.text = "DevOps & MLOps Orchestration Pipeline Dashboard"
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    st_box = slide1.shapes.add_textbox(Inches(1.2), Inches(2.6), Inches(10.9), Inches(0.6))
    st_frame = st_box.text_frame
    st_frame.word_wrap = True
    p = st_frame.paragraphs[0]
    p.text = "Unified Control Center for Infrastructure Telemetry, CI/CD Automation & ML Lifecycle Management"
    p.font.size = Pt(15)
    p.font.color.rgb = TEXT_MUTED

    meta_info = [
        ("Team Lead & Architect", "Boddiga Sai Krishna"),
        ("Frontend & UI/UX Developer", "Nagamaheshwar Reddy"),
        ("Backend & MLOps Specialist", "Raghu Lakshman"),
        ("Academic Review Period", "Capstone Review – 1 (August 2026)")
    ]

    card_w = Inches(2.55)
    card_h = Inches(1.8)
    start_x = Inches(1.2)
    gap_x = Inches(0.24)
    top_y = Inches(4.3)

    for i, (label, val) in enumerate(meta_info):
        cx = start_x + i * (card_w + gap_x)
        c_shape = add_card(slide1, cx, top_y, card_w, card_h, BG_COLOR, ACCENT_CYAN if i==0 else CARD_BORDER)
        
        tb = slide1.shapes.add_textbox(cx + Inches(0.15), top_y + Inches(0.2), card_w - Inches(0.3), card_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p1 = tf.paragraphs[0]
        p1.text = label.upper()
        p1.font.size = Pt(10)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_CYAN if i==0 else TEXT_MUTED
        
        p2 = tf.add_paragraph()
        p2.text = val
        p2.font.size = Pt(13)
        p2.font.bold = True
        p2.font.color.rgb = TEXT_WHITE
        p2.space_before = Pt(8)

    # ==========================================
    # SLIDE 2: Introduction & Motivation
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Introduction & Operational Motivation")

    b_tb = slide2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.5))
    p = b_tb.text_frame.paragraphs[0]
    p.text = "Modern development teams operate across separate, siloed administration consoles."
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    c_card = add_card(slide2, Inches(0.8), Inches(2.2), Inches(5.6), Inches(4.5))
    c_tb = slide2.shapes.add_textbox(Inches(1.0), Inches(2.4), Inches(5.2), Inches(4.1))
    c_tf = c_tb.text_frame
    c_tf.word_wrap = True

    p = c_tf.paragraphs[0]
    p.text = "THE OPERATIONAL CHALLENGE"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    challenges = [
        ("Tool Fragmentation", "GitHub for source control, Jenkins for CI/CD builds, Docker/Kubernetes for containers, and MLflow for ML tracking."),
        ("Cognitive Overhead", "Constant context-switching across browser tabs degrades incident response times and delays build visibility."),
        ("CORS Integration Barriers", "Cross-Origin Resource Sharing (CORS) blocks local browser applications from triggering internal CI/CD automation APIs.")
    ]

    for title, desc in challenges:
        pt = c_tf.add_paragraph()
        pt.text = "• " + title
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = TEXT_WHITE
        pt.space_before = Pt(10)

        pd = c_tf.add_paragraph()
        pd.text = desc
        pd.font.size = Pt(11)
        pd.font.color.rgb = TEXT_MUTED
        pd.space_before = Pt(2)

    # Right: Proposed Solution Banner
    s_card = add_card(slide2, Inches(6.8), Inches(2.2), Inches(5.733), Inches(4.5), CARD_BG, ACCENT_INDIGO)
    s_tb = slide2.shapes.add_textbox(Inches(7.0), Inches(2.4), Inches(5.333), Inches(4.1))
    s_tf = s_tb.text_frame
    s_tf.word_wrap = True

    p = s_tf.paragraphs[0]
    p.text = "THE UNIFIED SOLUTION"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO

    p_sol = s_tf.add_paragraph()
    p_sol.text = "Engineer a single-tenant React 18 dashboard combining live Git repository management, Jenkins build orchestration via Vite Reverse Proxy, and real-time 2.5s telemetry streams."
    p_sol.font.size = Pt(14)
    p_sol.font.color.rgb = TEXT_WHITE
    p_sol.space_before = Pt(14)

    steps = [
        "1. Live GitHub REST API v3 Integration",
        "2. CORS-Bypassed Jenkins CI/CD Proxy (/jenkins-proxy)",
        "3. Real-Time Telemetry Stream (CPU, RAM, Latency)",
        "4. MLOps Model Version & Accuracy Matrix"
    ]

    for s in steps:
        ps = s_tf.add_paragraph()
        ps.text = "✓ " + s
        ps.font.size = Pt(12)
        ps.font.bold = True
        ps.font.color.rgb = ACCENT_GREEN
        ps.space_before = Pt(10)

    # ==========================================
    # SLIDE 3: Problem Statement
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "Problem Statement")

    prob_cards = [
        ("Effort & Context Switching", "Navigating between GitHub, Jenkins UI, Grafana, and MLflow creates context-switching fatigue."),
        ("Documentation & Build Drift", "Build statuses and ML model metrics become desynchronized with active repository code."),
        ("CORS Security Restrictions", "Browser cross-origin security rules block direct REST API triggering of internal Jenkins build servers."),
        ("High APM Tool Costs", "Existing enterprise APM tools (Grafana/Datadog) are complex and expensive for single-tenant control."),
        ("Lack of MLOps Integration", "Traditional CI/CD tools lack native tracking for machine learning model training parameters and accuracy."),
        ("Delayed Visibility", "Lack of unified real-time telemetry delays incident response and pipeline build feedback.")
    ]

    p_w = Inches(3.65)
    p_h = Inches(1.8)

    for i, (title, desc) in enumerate(prob_cards):
        row = i // 2
        col = i % 2
        cx = Inches(0.8) + col * Inches(3.9)
        cy = Inches(1.6) + row * Inches(1.9)

        add_card(slide3, cx, cy, p_w, p_h)
        tb = slide3.shapes.add_textbox(cx + Inches(0.15), cy + Inches(0.15), p_w - Inches(0.3), p_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_CYAN

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.space_before = Pt(4)

    # Right side Quote callout
    q_card = add_card(slide3, Inches(8.8), Inches(1.6), Inches(3.733), Inches(5.4), CARD_BG, ACCENT_INDIGO)
    q_tb = slide3.shapes.add_textbox(Inches(9.0), Inches(2.2), Inches(3.333), Inches(4.2))
    q_tf = q_tb.text_frame
    q_tf.word_wrap = True

    p = q_tf.paragraphs[0]
    p.text = "FORMAL PROBLEM STATEMENT"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO

    p_q = q_tf.add_paragraph()
    p_q.text = "“Traditional DevOps platforms lack unified integration with machine learning lifecycles, causing fragmented telemetry, delayed build feedback, and CORS cross-origin integration barriers.”"
    p_q.font.size = Pt(16)
    p_q.font.bold = True
    p_q.font.italic = True
    p_q.font.color.rgb = TEXT_WHITE
    p_q.space_before = Pt(20)

    # ==========================================
    # SLIDE 4: Project Objectives
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Project Objectives")

    objs = [
        ("🤖 Centralized Control Center", "Develop a unified single-tenant web dashboard combining live GitHub repository management, Jenkins CI/CD execution, and ML model tracking."),
        ("🛡️ CORS-Bypassed Reverse Proxy", "Establish a Vite reverse proxy (`/jenkins-proxy`) coupled with ngrok tunneling to eliminate cross-origin security blocks when triggering Jenkins APIs."),
        ("📊 Real-Time Telemetry Stream", "Implement real-time 2.5-second refresh streams for CPU utilization, RAM consumption, and API response latency using Recharts."),
        ("🧠 MLOps Lifecycle Tracking", "Integrate Supabase PostgreSQL to persist machine learning model versions, training parameters, accuracy metrics, and deployment endpoints."),
        ("⚡ Live Build & Log Streaming", "Provide one-click 'Build Now' pipeline triggers and stream live Jenkins build console logs directly into the browser terminal interface."),
        ("🚀 Sub-200ms Latency & Modular Design", "Ensure responsive performance (<200ms API latency) with a modular architecture segregating `/frontend` (React 18) and `/backend` (Supabase).")
    ]

    o_w = Inches(3.65)
    o_h = Inches(2.4)

    for i, (title, desc) in enumerate(objs):
        col = i // 2
        row = i % 2
        cx = Inches(0.8) + col * Inches(3.9)
        cy = Inches(1.6) + row * Inches(2.6)

        add_card(slide4, cx, cy, o_w, o_h)
        tb = slide4.shapes.add_textbox(cx + Inches(0.15), cy + Inches(0.15), o_w - Inches(0.3), o_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_CYAN

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_WHITE
        p_d.space_before = Pt(6)

    # Right Callout Card
    r_card = add_card(slide4, Inches(8.8), Inches(1.6), Inches(3.733), Inches(5.0), CARD_BG, ACCENT_GREEN)
    r_tb = slide4.shapes.add_textbox(Inches(9.0), Inches(1.8), Inches(3.333), Inches(4.6))
    r_tf = r_tb.text_frame
    r_tf.word_wrap = True

    p = r_tf.paragraphs[0]
    p.text = "CORE MISSION"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    p_m = r_tf.add_paragraph()
    p_m.text = "Bridge traditional continuous delivery pipelines with machine learning operations under one unified, high-performance interface."
    p_m.font.size = Pt(15)
    p_m.font.bold = True
    p_m.font.color.rgb = TEXT_WHITE
    p_m.space_before = Pt(14)

    p_list = [
        "Eliminate context switching",
        "Provide instantaneous build feedback",
        "Stream live system telemetry",
        "Ensure enterprise-grade security"
    ]

    for item in p_list:
        pi = r_tf.add_paragraph()
        pi.text = "✓ " + item
        pi.font.size = Pt(12)
        pi.font.color.rgb = TEXT_MUTED
        pi.space_before = Pt(10)

    # ==========================================
    # SLIDE 5: Proposed System Pipeline
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Proposed System Pipeline")

    b_tb = slide5.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.4))
    p = b_tb.text_frame.paragraphs[0]
    p.text = "The platform unifies developer repository events, Jenkins build execution, and MLOps metrics into one flow."
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE

    steps_data = [
        ("1. GitHub Event Trigger", "User authenticates via PAT; live commits, branches, and repo changes are monitored via API v3."),
        ("2. Vite Reverse Proxy (/jenkins-proxy)", "Browser dispatches build triggers bypassing CORS restriction to internal Jenkins server (:8080)."),
        ("3. Jenkins Build Orchestration", "Jenkins executes build pipelines ('Skill 7', 'Maven-Build-Job') and streams live console logs to browser."),
        ("4. Real-Time Telemetry & MLOps Matrix", "Recharts streams 2.5s CPU/RAM/Latency metrics while Supabase persists ML model accuracy logs.")
    ]

    p_w = Inches(11.733)
    p_h = Inches(1.1)

    for i, (title, desc) in enumerate(steps_data):
        cy = Inches(2.1) + i * Inches(1.2)
        c_shape = add_card(slide5, Inches(0.8), cy, p_w, p_h, CARD_BG if i%2==0 else BG_COLOR, ACCENT_CYAN)

        tb = slide5.shapes.add_textbox(Inches(1.1), cy + Inches(0.15), p_w - Inches(0.6), p_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_CYAN

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(12)
        p_d.font.color.rgb = TEXT_WHITE
        p_d.space_before = Pt(2)

    # ==========================================
    # SLIDE 6: System Architecture
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "System Architecture & Layered Integration")

    arch_layers = [
        ("CLIENT LAYER (Browser :5173)", "React 18 + TypeScript UI  |  Recharts Real-Time Telemetry Stream", "Renders dark glassmorphism dashboard, manages state, and displays live build console streams.", ACCENT_CYAN),
        ("PROXY & GATEWAY LAYER", "Vite Reverse Proxy (/jenkins-proxy)  ──►  Localhost / ngrok Tunnel", "Bypasses browser CORS cross-origin policies by forwarding API calls to internal Jenkins build server.", ACCENT_INDIGO),
        ("INTEGRATIONS & STORAGE LAYER", "GitHub REST API v3  |  Jenkins CI/CD Server  |  Supabase PostgreSQL DB", "Handles persistent storage of MLOps model versions, GitHub commit logs, and live build execution triggers.", ACCENT_GREEN)
    ]

    top_pos = Inches(1.6)
    layer_h = Inches(1.6)
    gap_y = Inches(0.2)

    for i, (title, sub, desc, accent) in enumerate(arch_layers):
        cy = top_pos + i * (layer_h + gap_y)
        card = add_card(slide6, Inches(0.8), cy, Inches(11.733), layer_h, CARD_BG, accent)

        tb = slide6.shapes.add_textbox(Inches(1.1), cy + Inches(0.15), Inches(11.1), layer_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = f"LAYER {i+1}: {title}"
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = accent

        p_sub = tf.add_paragraph()
        p_sub.text = sub
        p_sub.font.size = Pt(13)
        p_sub.font.bold = True
        p_sub.font.color.rgb = TEXT_WHITE
        p_sub.space_before = Pt(4)

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.space_before = Pt(4)

    # ==========================================
    # SLIDE 7: Technology Stack
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Technology Stack Specifications")

    tech_categories = [
        ("Frontend Framework", "React v18.3.1 & TypeScript v5.5.3"),
        ("Build Tool & Bundler", "Vite v5.4.8 (Reverse Proxy Configuration Enabled)"),
        ("Styling & Icons", "Tailwind CSS v3.4.1 & Lucide React v0.344"),
        ("Data Visualization", "Recharts v3.10.0 (Real-Time Infrastructure Telemetry)"),
        ("Backend & Database", "Supabase PostgreSQL BaaS & @supabase/supabase-js v2.57"),
        ("CI/CD Automation", "Jenkins v2.528.1 & GitHub REST API v3"),
        ("Local Proxy & Tunnel", "Vite Reverse Proxy (/jenkins-proxy) & ngrok Tunnels"),
        ("Linting & Code Quality", "ESLint v9.9.1 & Modular Architecture Segregation")
    ]

    t_w = Inches(5.7)
    t_h = Inches(1.2)

    for i, (cat, stack) in enumerate(tech_categories):
        col = i // 4
        row = i % 4
        cx = Inches(0.8) + col * Inches(6.033)
        cy = Inches(1.6) + row * Inches(1.35)

        add_card(slide7, cx, cy, t_w, t_h)
        tb = slide7.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.15), t_w - Inches(0.4), t_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p_c = tf.paragraphs[0]
        p_c.text = cat
        p_c.font.size = Pt(13)
        p_c.font.bold = True
        p_c.font.color.rgb = ACCENT_CYAN

        p_s = tf.add_paragraph()
        p_s.text = stack
        p_s.font.size = Pt(12)
        p_s.font.color.rgb = TEXT_WHITE
        p_s.space_before = Pt(4)

    # ==========================================
    # SLIDE 8: Easy Team Task Distribution (3 Cards)
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Team Task Distribution & Member Focus")

    team_allocations = [
        ("Boddiga Sai Krishna", "Team Lead & Full-Stack Architect", "System Architecture & API Integration", [
            "Project Leadership & System Architecture",
            "Vite Reverse Proxy setup (/jenkins-proxy)",
            "GitHub REST API v3 & Jenkins REST Integration",
            "CORS Cross-Origin Security Bypass Engine"
        ], ACCENT_CYAN),
        ("Nagamaheshwar Reddy", "UI/UX & Frontend Developer", "React Dashboard & Telemetry Visualizations", [
            "React 18 Component Architecture",
            "Tailwind CSS Glassmorphism Design System",
            "Recharts Infrastructure Telemetry Engine",
            "Live Jenkins Terminal & Build Inspector UI"
        ], ACCENT_INDIGO),
        ("Raghu Lakshman", "MLOps & Database Specialist", "Supabase Database & ML Model Matrix", [
            "Supabase PostgreSQL Schema & Migrations",
            "ML Model Version & Accuracy Matrix",
            "Microservice Component Health Matrix",
            "Quality Assurance, Unit Testing & Audit Logs"
        ], ACCENT_GREEN)
    ]

    m_w = Inches(3.65)
    m_h = Inches(5.2)
    m_start_x = Inches(0.8)
    m_gap = Inches(0.39)

    for i, (name, role, summary, tasks, accent) in enumerate(team_allocations):
        cx = m_start_x + i * (m_w + m_gap)
        c_shape = add_card(slide8, cx, Inches(1.6), m_w, m_h, CARD_BG, accent)

        tb = slide8.shapes.add_textbox(cx + Inches(0.15), Inches(1.8), m_w - Inches(0.3), m_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p_name = tf.paragraphs[0]
        p_name.text = name
        p_name.font.size = Pt(16)
        p_name.font.bold = True
        p_name.font.color.rgb = accent

        p_role = tf.add_paragraph()
        p_role.text = role
        p_role.font.size = Pt(12)
        p_role.font.bold = True
        p_role.font.color.rgb = TEXT_WHITE
        p_role.space_before = Pt(4)

        p_sum = tf.add_paragraph()
        p_sum.text = f"Primary Focus: {summary}"
        p_sum.font.size = Pt(10)
        p_sum.font.bold = True
        p_sum.font.color.rgb = TEXT_MUTED
        p_sum.space_before = Pt(6)

        p_head = tf.add_paragraph()
        p_head.text = "KEY DELIVERABLES:"
        p_head.font.size = Pt(11)
        p_head.font.bold = True
        p_head.font.color.rgb = accent
        p_head.space_before = Pt(14)

        for t in tasks:
            pt = tf.add_paragraph()
            pt.text = "• " + t
            pt.font.size = Pt(11)
            pt.font.color.rgb = TEXT_WHITE
            pt.space_before = Pt(6)

    # ==========================================
    # SLIDE 9: Expected Outcomes
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "Expected Outcomes & Key Deliverables")

    outcomes = [
        ("⚡ Unified Single-Tenant Control Center", "Consolidates GitHub, Jenkins UI, Grafana, and MLflow into one dark-themed React dashboard."),
        ("🎯 100% CORS-Bypassed Proxy Routing", "Seamless remote triggering of internal Jenkins builds without cross-origin browser security blocks."),
        ("📊 Real-Time Infrastructure Observability", "Continuous 2.5-second telemetry streams for CPU, RAM, and Latency accompanied by component health matrix."),
        ("🧠 Integrated MLOps Lifecycle Matrix", "Native tracking for ML model versions, framework parameters, training accuracy, and deployment endpoints."),
        ("⚡ High-Performance Sub-200ms Latency", "Lightweight, responsive UI built with Vite, React 18, and segregated modular frontend/backend architecture.")
    ]

    for i, (title, desc) in enumerate(outcomes):
        cy = Inches(1.6) + i * Inches(1.05)
        add_card(slide9, Inches(0.8), cy, Inches(11.733), Inches(0.95), CARD_BG, ACCENT_GREEN)

        tb = slide9.shapes.add_textbox(Inches(1.1), cy + Inches(0.1), Inches(11.133), Inches(0.75))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_GREEN

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_WHITE
        p_d.space_before = Pt(2)

    # ==========================================
    # SLIDE 10: Project Roadmap & Future Scope
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)
    add_header(slide10, "Project Roadmap & Future Scope")

    futures = [
        ("PHASE 1 (Aug 17–22, 2026) - COMPLETED", "Capstone Review-1 Initiation, Problem Definition, Architecture, Live GitHub & Jenkins Proxy setup."),
        ("PHASE 2 (Sept 2026) - CONTAINERIZATION", "Docker Engine Integration, Container metrics monitoring & Kubernetes Cluster health views."),
        ("PHASE 3 (Oct 2026) - CLOUD & MLOPS", "AWS EC2/S3 Cloud integration & Automated ML Model deployment pipelines."),
        ("PHASE 4 (Nov 2026) - AUDIT & FINAL", "Final System Audit, Security Penetration Testing, Capstone Final Review & Cloud Deployment."),
        ("AI Incident Diagnostics", "Integration of LLM agents for automated build failure root-cause analysis and log troubleshooting.")
    ]

    for i, (title, desc) in enumerate(futures):
        cy = Inches(1.6) + i * Inches(1.05)
        add_card(slide10, Inches(0.8), cy, Inches(11.733), Inches(0.95))

        tb = slide10.shapes.add_textbox(Inches(1.1), cy + Inches(0.1), Inches(11.133), Inches(0.75))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = ACCENT_CYAN

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_WHITE
        p_d.space_before = Pt(2)

    # ==========================================
    # SLIDE 11: Conclusion
    # ==========================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide11)
    add_header(slide11, "Conclusion")

    add_card(slide11, Inches(0.8), Inches(1.6), Inches(11.733), Inches(5.2), CARD_BG, ACCENT_CYAN)

    tb = slide11.shapes.add_textbox(Inches(1.2), Inches(2.0), Inches(10.933), Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "SUMMARY OF PROJECT ACHIEVEMENTS"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    c_points = [
        ("Bridged DevOps and MLOps Divide", "The platform successfully eliminates operational silos by bringing continuous integration build controls and machine learning lifecycle management into one interface."),
        ("Solved CORS Integration Barriers", "Engineered a robust Vite reverse proxy mechanism coupled with ngrok tunneling to allow secure browser interaction with local Jenkins servers."),
        ("Real-Time Telemetry Observability", "Integrated Recharts to stream infrastructure CPU, RAM, and latency metrics at 2.5-second intervals with microservice health monitoring."),
        ("Production-Ready Full-Stack Architecture", "Segregated modular frontend (React 18 + TypeScript + Vite) and backend (Supabase PostgreSQL) ensures scalable enterprise deployment.")
    ]

    for title, desc in c_points:
        p_t = tf.add_paragraph()
        p_t.text = "• " + title
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = TEXT_WHITE
        p_t.space_before = Pt(12)

        p_d = tf.add_paragraph()
        p_d.text = "   " + desc
        p_d.font.size = Pt(11)
        p_d.font.color.rgb = TEXT_MUTED
        p_d.space_before = Pt(2)

    # ==========================================
    # SLIDE 12: Thank You / Q&A Slide
    # ==========================================
    slide12 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide12)

    add_card(slide12, Inches(1.5), Inches(1.5), Inches(10.333), Inches(4.5), CARD_BG, ACCENT_CYAN)
    tb = slide12.shapes.add_textbox(Inches(1.8), Inches(2.0), Inches(9.733), Inches(3.5))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Thank You!"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN

    p_sub = tf.add_paragraph()
    p_sub.text = "DevOps & MLOps Orchestration Pipeline Dashboard"
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.font.size = Pt(18)
    p_sub.font.bold = True
    p_sub.font.color.rgb = TEXT_WHITE
    p_sub.space_before = Pt(8)

    p_team = tf.add_paragraph()
    p_team.text = "Boddiga Sai Krishna (Lead)  |  Nagamaheshwar Reddy  |  Raghu Lakshman"
    p_team.alignment = PP_ALIGN.CENTER
    p_team.font.size = Pt(14)
    p_team.font.bold = True
    p_team.font.color.rgb = ACCENT_INDIGO
    p_team.space_before = Pt(12)

    p_qa = tf.add_paragraph()
    p_qa.text = "Capstone Review – 1  |  Questions & Answers / Discussion"
    p_qa.alignment = PP_ALIGN.CENTER
    p_qa.font.size = Pt(13)
    p_qa.font.color.rgb = TEXT_MUTED
    p_qa.space_before = Pt(12)

    output_path = r"c:\Users\bsaik\Desktop\Capstone\pipeline-main\DevOps_MLOps_Pipeline_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    build_devops_mlops_presentation()
